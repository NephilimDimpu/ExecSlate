"""
ExecSlate — Analytics Workspace PDF Exporter
A lightweight PDF brief generated directly from a saved Analytics session.
Focused on the exploratory output (KPIs + charts + observations + notes),
not the full board-ready report (that's exports/export_functions.py).
"""

import base64
import io
import logging
from datetime import datetime
from pathlib import Path

from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Image, Table, TableStyle, PageBreak
)
from reportlab.lib.units import inch
from reportlab.lib.enums import TA_LEFT, TA_CENTER
from reportlab.lib import colors

logger = logging.getLogger("ExecSlate")


def _fmt_value(value, fmt, currency="$"):
    """Format a KPI value based on its declared format type."""
    try:
        v = float(value)
    except Exception:
        return str(value)
    if fmt == "currency":
        if abs(v) >= 1_000_000:
            return f"{currency}{v/1_000_000:.2f}M"
        if abs(v) >= 1_000:
            return f"{currency}{v/1_000:.1f}K"
        return f"{currency}{v:,.0f}"
    if fmt == "percent":
        return f"{v:.1f}%"
    if abs(v) >= 1_000_000:
        return f"{v/1_000_000:.2f}M"
    if abs(v) >= 1_000:
        return f"{v:,.0f}"
    return f"{v:,.0f}"


def _b64_to_image_flowable(b64_string, max_width_inches=6.5):
    """Decode a base64 PNG into a ReportLab Image flowable. Returns None on failure."""
    if not b64_string:
        return None
    try:
        raw = base64.b64decode(b64_string)
        bio = io.BytesIO(raw)
        img = Image(bio, width=max_width_inches * inch, height=(max_width_inches * 0.55) * inch)
        img.hAlign = 'CENTER'
        return img
    except Exception as e:
        logger.warning(f"Could not embed chart image: {e}")
        return None


def export_analytics_pdf(project, session, latest_draft, output_path):
    """
    Build a one-page-style PDF brief for the Analytics workspace.

    Args:
        project: project dict (client, period, etc.)
        session: latest analytics session dict (kpi_metrics, ai_insights, charts, ...)
        latest_draft: optional saved-draft dict with notes + selected_insights
        output_path: where to write the PDF
    """
    output_path = Path(output_path)
    doc = SimpleDocTemplate(
        str(output_path), pagesize=A4,
        leftMargin=0.6 * inch, rightMargin=0.6 * inch,
        topMargin=0.5 * inch, bottomMargin=0.5 * inch,
        title=f"ExecSlate Analytics — {project.get('client', '')}",
    )

    styles = getSampleStyleSheet()
    brand = ParagraphStyle("Brand", parent=styles["Normal"],
                           fontName="Helvetica-Bold", fontSize=9,
                           textColor=colors.HexColor("#1e40af"),
                           alignment=TA_LEFT, spaceAfter=2)
    title = ParagraphStyle("Title", parent=styles["Title"], fontSize=20,
                           textColor=colors.HexColor("#0f172a"),
                           spaceAfter=4, alignment=TA_LEFT)
    sub = ParagraphStyle("Sub", parent=styles["Normal"], fontSize=10,
                         textColor=colors.HexColor("#64748b"), spaceAfter=14)
    h2 = ParagraphStyle("H2", parent=styles["Heading2"], fontSize=12,
                        textColor=colors.HexColor("#0f172a"),
                        spaceBefore=14, spaceAfter=8)
    body = ParagraphStyle("Body", parent=styles["Normal"], fontSize=10,
                          textColor=colors.HexColor("#1e293b"),
                          leading=15, spaceAfter=5)
    bullet = ParagraphStyle("Bullet", parent=body, leftIndent=14,
                            bulletIndent=2, spaceAfter=4)
    note = ParagraphStyle("Note", parent=body,
                          textColor=colors.HexColor("#475569"),
                          fontName="Helvetica-Oblique")

    story = []
    story.append(Paragraph("EXECSLATE — ANALYTICS BRIEF", brand))
    story.append(Paragraph(project.get("client") or "Untitled Project", title))
    period = project.get("period") or ""
    generated = datetime.now().strftime("%B %d, %Y")
    sub_text = f"Analytics workspace export &nbsp;·&nbsp; {period} &nbsp;·&nbsp; Generated {generated}"
    story.append(Paragraph(sub_text, sub))

    # ── Stat bar ──
    currency = session.get("currency") or project.get("currency") or "$"
    stat_cells = []
    if session.get("row_count"):
        stat_cells.append([Paragraph("<b>Rows</b>", body), Paragraph(f"{session['row_count']:,}", body)])
    if session.get("confidence"):
        stat_cells.append([Paragraph("<b>Confidence</b>", body), Paragraph(f"{session['confidence']}%", body)])
    if session.get("trend"):
        stat_cells.append([Paragraph("<b>Pattern</b>", body), Paragraph(session["trend"], body)])
    if session.get("growth_rate"):
        stat_cells.append([Paragraph("<b>Growth</b>", body),
                           Paragraph(f"{session['growth_rate']:.1f}%", body)])
    if session.get("primary_col"):
        primary_label = str(session["primary_col"]).replace("_", " ").title()
        stat_cells.append([Paragraph("<b>Primary KPI</b>", body), Paragraph(primary_label, body)])
    if stat_cells:
        stat_table = Table(stat_cells, colWidths=[1.4 * inch, 5.4 * inch])
        stat_table.setStyle(TableStyle([
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("LINEBELOW", (0, 0), (-1, -1), 0.25, colors.HexColor("#e2e8f0")),
        ]))
        story.append(stat_table)
        story.append(Spacer(1, 6))

    # ── KPIs ──
    kpis = session.get("kpi_metrics") or {}
    if kpis:
        story.append(Paragraph("Key Performance Indicators", h2))
        kpi_rows = [["Metric", "Total", "Average", "Direction"]]
        for key, m in list(kpis.items())[:10]:
            label = m.get("label") or key
            fmt = m.get("format", "number")
            val = _fmt_value(m.get("value"), fmt, currency) if "value" in m else _fmt_value(m.get("total", 0), fmt, currency)
            avg = _fmt_value(m.get("avg", 0), fmt, currency)
            direction = "↑ good" if m.get("direction") == "up_good" else ("↓ good" if m.get("direction") == "down_good" else "—")
            kpi_rows.append([label, val, avg, direction])
        kpi_table = Table(kpi_rows, colWidths=[2.5 * inch, 1.6 * inch, 1.6 * inch, 1.1 * inch])
        kpi_table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1e3a8a")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 9),
            ("ALIGN", (1, 1), (-1, -1), "RIGHT"),
            ("ALIGN", (0, 0), (-1, 0), "LEFT"),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
            ("TOPPADDING", (0, 0), (-1, -1), 6),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1),
             [colors.white, colors.HexColor("#f8fafc")]),
            ("LINEBELOW", (0, 0), (-1, -1), 0.25, colors.HexColor("#e2e8f0")),
        ]))
        story.append(kpi_table)

    # ── Primary chart ──
    primary_img = _b64_to_image_flowable(session.get("revenue_chart_b64"))
    if primary_img:
        story.append(Spacer(1, 6))
        story.append(Paragraph("Primary View", h2))
        story.append(primary_img)

    # ── Breakdown chart ──
    breakdown_img = _b64_to_image_flowable(session.get("region_chart_b64"))
    if breakdown_img:
        story.append(Paragraph("Breakdown", h2))
        story.append(breakdown_img)

    # ── Observations ──
    insights = session.get("ai_insights") or []
    if insights:
        story.append(Paragraph("Observations", h2))
        for ins in insights:
            story.append(Paragraph(f"• {ins}", bullet))

    # ── Pinned + analyst notes from latest saved draft ──
    if latest_draft:
        pinned = latest_draft.get("selected_insights") or []
        notes_text = (latest_draft.get("notes") or "").strip()
        if pinned:
            story.append(Paragraph("Pinned for Decision", h2))
            for p in pinned:
                story.append(Paragraph(f"• {p}", bullet))
        if notes_text:
            story.append(Paragraph("Analyst Notes", h2))
            for line in notes_text.split("\n"):
                if line.strip():
                    story.append(Paragraph(line.strip(), body))

    # ── Footer ──
    story.append(Spacer(1, 18))
    footer = ParagraphStyle("Footer", parent=body, fontSize=8,
                            textColor=colors.HexColor("#94a3b8"),
                            alignment=TA_CENTER)
    story.append(Paragraph(
        f"Generated by ExecSlate · Analytics Workspace Export · {generated}",
        footer))

    doc.build(story)
    return output_path


# ════════════════════════════════════════════════════════════════════════════
# PPT EXPORT
# ════════════════════════════════════════════════════════════════════════════

def _decode_b64_to_bytesio(b64_string):
    if not b64_string:
        return None
    try:
        return io.BytesIO(base64.b64decode(b64_string))
    except Exception as e:
        logger.warning(f"Chart decode failed: {e}")
        return None


def export_analytics_pptx(project, session, latest_draft, output_path):
    """Build a slide deck export of the Analytics workspace."""
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN

    output_path = Path(output_path)
    prs = Presentation()
    prs.slide_width = Inches(13.33)
    prs.slide_height = Inches(7.5)

    BLANK = prs.slide_layouts[6]
    NAVY = RGBColor(0x0F, 0x17, 0x2A)
    BLUE = RGBColor(0x1E, 0x40, 0xAF)
    MUTED = RGBColor(0x64, 0x74, 0x8B)
    WHITE = RGBColor(0xFF, 0xFF, 0xFF)

    def add_title(slide, text, size=28, color=NAVY, top=0.4):
        box = slide.shapes.add_textbox(Inches(0.6), Inches(top), Inches(12), Inches(0.9))
        tf = box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = text
        p.font.size = Pt(size)
        p.font.bold = True
        p.font.color.rgb = color

    def add_subtitle(slide, text, top=1.1, size=12, color=MUTED):
        box = slide.shapes.add_textbox(Inches(0.6), Inches(top), Inches(12), Inches(0.6))
        tf = box.text_frame
        p = tf.paragraphs[0]
        p.text = text
        p.font.size = Pt(size)
        p.font.color.rgb = color

    # ── Title slide ──
    s = prs.slides.add_slide(BLANK)
    bg = s.shapes.add_shape(1, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = NAVY
    bg.line.fill.background()
    tb = s.shapes.add_textbox(Inches(0.8), Inches(2.5), Inches(11.7), Inches(1.2))
    tb.text_frame.paragraphs[0].text = "ExecSlate — Analytics Brief"
    tb.text_frame.paragraphs[0].font.size = Pt(14)
    tb.text_frame.paragraphs[0].font.color.rgb = RGBColor(0x93, 0xC5, 0xFD)
    tb.text_frame.paragraphs[0].font.bold = True
    tb2 = s.shapes.add_textbox(Inches(0.8), Inches(3.0), Inches(11.7), Inches(2.0))
    p2 = tb2.text_frame.paragraphs[0]
    p2.text = project.get("client") or "Untitled Project"
    p2.font.size = Pt(44)
    p2.font.bold = True
    p2.font.color.rgb = WHITE
    period = project.get("period") or ""
    generated = datetime.now().strftime("%B %d, %Y")
    tb3 = s.shapes.add_textbox(Inches(0.8), Inches(5.0), Inches(11.7), Inches(0.6))
    p3 = tb3.text_frame.paragraphs[0]
    p3.text = f"{period}    ·    Generated {generated}"
    p3.font.size = Pt(14)
    p3.font.color.rgb = RGBColor(0xCB, 0xD5, 0xE1)

    # ── KPI slide ──
    currency = session.get("currency") or project.get("currency") or "$"
    kpis = session.get("kpi_metrics") or {}
    if kpis:
        s = prs.slides.add_slide(BLANK)
        add_title(s, "Key Performance Indicators")
        add_subtitle(s, f"{session.get('row_count', 0):,} rows · Confidence {session.get('confidence', 0)}% · Pattern: {session.get('trend', '—')}")
        items = list(kpis.items())[:6]
        cols = min(3, len(items))
        rows = (len(items) + cols - 1) // cols
        card_w = (12.0 - (cols - 1) * 0.3) / cols
        card_h = 1.8
        for idx, (key, m) in enumerate(items):
            r, c = divmod(idx, cols)
            x = 0.6 + c * (card_w + 0.3)
            y = 1.9 + r * (card_h + 0.3)
            card = s.shapes.add_shape(5, Inches(x), Inches(y), Inches(card_w), Inches(card_h))
            card.fill.solid()
            card.fill.fore_color.rgb = WHITE
            card.line.color.rgb = RGBColor(0xE2, 0xE8, 0xF0)
            tb = s.shapes.add_textbox(Inches(x + 0.2), Inches(y + 0.2), Inches(card_w - 0.4), Inches(card_h - 0.4))
            tf = tb.text_frame
            tf.word_wrap = True
            label_p = tf.paragraphs[0]
            label_p.text = (m.get("label") or key).upper()
            label_p.font.size = Pt(10)
            label_p.font.bold = True
            label_p.font.color.rgb = MUTED
            val_p = tf.add_paragraph()
            fmt = m.get("format", "number")
            val_p.text = _fmt_value(m.get("value", m.get("total", 0)), fmt, currency)
            val_p.font.size = Pt(22)
            val_p.font.bold = True
            val_p.font.color.rgb = NAVY
            avg_p = tf.add_paragraph()
            avg_p.text = f"avg {_fmt_value(m.get('avg', 0), fmt, currency)}"
            avg_p.font.size = Pt(10)
            avg_p.font.color.rgb = MUTED

    # ── Primary chart slide ──
    primary_io = _decode_b64_to_bytesio(session.get("revenue_chart_b64"))
    if primary_io:
        s = prs.slides.add_slide(BLANK)
        add_title(s, "Primary View")
        s.shapes.add_picture(primary_io, Inches(1.2), Inches(1.4), width=Inches(11), height=Inches(5.6))

    # ── Breakdown chart slide ──
    breakdown_io = _decode_b64_to_bytesio(session.get("region_chart_b64"))
    if breakdown_io:
        s = prs.slides.add_slide(BLANK)
        add_title(s, "Breakdown")
        s.shapes.add_picture(breakdown_io, Inches(1.2), Inches(1.4), width=Inches(11), height=Inches(5.6))

    # ── Observations slide ──
    insights = session.get("ai_insights") or []
    if insights:
        s = prs.slides.add_slide(BLANK)
        add_title(s, "Observations")
        tb = s.shapes.add_textbox(Inches(0.8), Inches(1.4), Inches(11.7), Inches(5.5))
        tf = tb.text_frame
        tf.word_wrap = True
        for i, ins in enumerate(insights):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = f"•  {ins}"
            p.font.size = Pt(15)
            p.font.color.rgb = NAVY
            p.space_after = Pt(8)

    # ── Pinned + notes slide ──
    if latest_draft:
        pinned = latest_draft.get("selected_insights") or []
        notes_text = (latest_draft.get("notes") or "").strip()
        if pinned or notes_text:
            s = prs.slides.add_slide(BLANK)
            add_title(s, "Pinned & Analyst Notes")
            tb = s.shapes.add_textbox(Inches(0.8), Inches(1.4), Inches(11.7), Inches(5.5))
            tf = tb.text_frame
            tf.word_wrap = True
            first = True
            if pinned:
                head = tf.paragraphs[0]
                head.text = "Pinned for decision:"
                head.font.size = Pt(15)
                head.font.bold = True
                head.font.color.rgb = BLUE
                first = False
                for ins in pinned:
                    p = tf.add_paragraph()
                    p.text = f"•  {ins}"
                    p.font.size = Pt(14)
                    p.font.color.rgb = NAVY
            if notes_text:
                p = tf.add_paragraph() if not first else tf.paragraphs[0]
                p.text = ("Analyst notes:" if not first else "Analyst notes:")
                p.font.size = Pt(15)
                p.font.bold = True
                p.font.color.rgb = BLUE
                p.space_before = Pt(14)
                for line in notes_text.split("\n"):
                    if line.strip():
                        np = tf.add_paragraph()
                        np.text = line.strip()
                        np.font.size = Pt(13)
                        np.font.color.rgb = NAVY

    prs.save(str(output_path))
    return output_path


# ════════════════════════════════════════════════════════════════════════════
# WORD (.docx) EXPORT
# ════════════════════════════════════════════════════════════════════════════

def export_analytics_docx(project, session, latest_draft, output_path):
    """Editable Word document mirroring the PDF brief."""
    from docx import Document
    from docx.shared import Inches as DocxInches, Pt as DocxPt, RGBColor as DocxRGB
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    output_path = Path(output_path)
    doc = Document()

    # Title
    h = doc.add_paragraph()
    h.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = h.add_run("EXECSLATE — ANALYTICS BRIEF")
    run.bold = True
    run.font.size = DocxPt(10)
    run.font.color.rgb = DocxRGB(0x1E, 0x40, 0xAF)

    t = doc.add_paragraph()
    tr = t.add_run(project.get("client") or "Untitled Project")
    tr.bold = True
    tr.font.size = DocxPt(22)
    tr.font.color.rgb = DocxRGB(0x0F, 0x17, 0x2A)

    period = project.get("period") or ""
    generated = datetime.now().strftime("%B %d, %Y")
    sub = doc.add_paragraph()
    sr = sub.add_run(f"Analytics workspace export  ·  {period}  ·  Generated {generated}")
    sr.italic = True
    sr.font.size = DocxPt(10)
    sr.font.color.rgb = DocxRGB(0x64, 0x74, 0x8B)

    # Stat bar
    bar = doc.add_paragraph()
    parts = []
    if session.get("row_count"):
        parts.append(f"{session['row_count']:,} rows")
    if session.get("confidence"):
        parts.append(f"Confidence {session['confidence']}%")
    if session.get("trend"):
        parts.append(f"Pattern: {session['trend']}")
    if session.get("growth_rate"):
        parts.append(f"Growth: {session['growth_rate']:.1f}%")
    if session.get("primary_col"):
        parts.append(f"Primary KPI: {str(session['primary_col']).replace('_', ' ').title()}")
    if parts:
        br = bar.add_run("   ·   ".join(parts))
        br.font.size = DocxPt(10)
        br.font.color.rgb = DocxRGB(0x47, 0x55, 0x69)

    # KPIs
    currency = session.get("currency") or project.get("currency") or "$"
    kpis = session.get("kpi_metrics") or {}
    if kpis:
        doc.add_heading("Key Performance Indicators", level=2)
        table = doc.add_table(rows=1, cols=4)
        table.style = "Light Grid Accent 1"
        hdr = table.rows[0].cells
        hdr[0].text = "Metric"
        hdr[1].text = "Total"
        hdr[2].text = "Average"
        hdr[3].text = "Direction"
        for key, m in list(kpis.items())[:10]:
            row = table.add_row().cells
            row[0].text = m.get("label") or key
            fmt = m.get("format", "number")
            row[1].text = _fmt_value(m.get("value", m.get("total", 0)), fmt, currency)
            row[2].text = _fmt_value(m.get("avg", 0), fmt, currency)
            row[3].text = ("↑ good" if m.get("direction") == "up_good"
                           else "↓ good" if m.get("direction") == "down_good" else "—")

    # Primary chart
    primary_io = _decode_b64_to_bytesio(session.get("revenue_chart_b64"))
    if primary_io:
        doc.add_heading("Primary View", level=2)
        doc.add_picture(primary_io, width=DocxInches(6.3))

    # Breakdown
    breakdown_io = _decode_b64_to_bytesio(session.get("region_chart_b64"))
    if breakdown_io:
        doc.add_heading("Breakdown", level=2)
        doc.add_picture(breakdown_io, width=DocxInches(6.3))

    # Observations
    insights = session.get("ai_insights") or []
    if insights:
        doc.add_heading("Observations", level=2)
        for ins in insights:
            p = doc.add_paragraph(style="List Bullet")
            p.add_run(ins)

    # Pinned + notes
    if latest_draft:
        pinned = latest_draft.get("selected_insights") or []
        notes_text = (latest_draft.get("notes") or "").strip()
        if pinned:
            doc.add_heading("Pinned for Decision", level=2)
            for p_text in pinned:
                doc.add_paragraph(p_text, style="List Bullet")
        if notes_text:
            doc.add_heading("Analyst Notes", level=2)
            for line in notes_text.split("\n"):
                if line.strip():
                    doc.add_paragraph(line.strip())

    doc.save(str(output_path))
    return output_path


# ════════════════════════════════════════════════════════════════════════════
# EXCEL (.xlsx) EXPORT
# ════════════════════════════════════════════════════════════════════════════

def export_analytics_xlsx(project, session, latest_draft, output_path):
    """Structured workbook: Summary / KPIs / Observations / Notes sheets."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    output_path = Path(output_path)
    wb = Workbook()

    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill("solid", fgColor="1E3A8A")
    bold_font = Font(bold=True, color="1E293B")

    # ── Summary sheet ──
    s1 = wb.active
    s1.title = "Summary"
    s1["A1"] = "ExecSlate — Analytics Brief"
    s1["A1"].font = Font(bold=True, size=14, color="0F172A")
    s1["A3"] = "Project"
    s1["A3"].font = bold_font
    s1["B3"] = project.get("client") or ""
    s1["A4"] = "Period"
    s1["A4"].font = bold_font
    s1["B4"] = project.get("period") or ""
    s1["A5"] = "Generated"
    s1["A5"].font = bold_font
    s1["B5"] = datetime.now().strftime("%Y-%m-%d %H:%M")
    s1["A7"] = "Rows"
    s1["A7"].font = bold_font
    s1["B7"] = session.get("row_count") or 0
    s1["A8"] = "Confidence (%)"
    s1["A8"].font = bold_font
    s1["B8"] = session.get("confidence") or 0
    s1["A9"] = "Pattern"
    s1["A9"].font = bold_font
    s1["B9"] = session.get("trend") or ""
    s1["A10"] = "Growth (%)"
    s1["A10"].font = bold_font
    s1["B10"] = round(session.get("growth_rate") or 0.0, 2)
    s1["A11"] = "Primary KPI"
    s1["A11"].font = bold_font
    s1["B11"] = (session.get("primary_col") or "").replace("_", " ").title()
    s1.column_dimensions["A"].width = 22
    s1.column_dimensions["B"].width = 40

    # ── KPIs sheet ──
    kpis = session.get("kpi_metrics") or {}
    if kpis:
        s2 = wb.create_sheet("KPIs")
        headers = ["Metric", "Total", "Average", "Format", "Direction"]
        for col, h in enumerate(headers, start=1):
            cell = s2.cell(row=1, column=col, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="left", vertical="center")
        for r, (key, m) in enumerate(kpis.items(), start=2):
            s2.cell(row=r, column=1, value=m.get("label") or key)
            s2.cell(row=r, column=2, value=float(m.get("value", m.get("total", 0)) or 0))
            s2.cell(row=r, column=3, value=float(m.get("avg", 0) or 0))
            s2.cell(row=r, column=4, value=m.get("format", "number"))
            s2.cell(row=r, column=5, value=m.get("direction", "neutral"))
        s2.column_dimensions["A"].width = 28
        for c in ("B", "C", "D", "E"):
            s2.column_dimensions[c].width = 16

    # ── Observations sheet ──
    insights = session.get("ai_insights") or []
    if insights:
        s3 = wb.create_sheet("Observations")
        s3.cell(row=1, column=1, value="#").font = header_font
        s3.cell(row=1, column=1).fill = header_fill
        s3.cell(row=1, column=2, value="Observation").font = header_font
        s3.cell(row=1, column=2).fill = header_fill
        for i, ins in enumerate(insights, start=1):
            s3.cell(row=i + 1, column=1, value=i)
            cell = s3.cell(row=i + 1, column=2, value=ins)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        s3.column_dimensions["A"].width = 6
        s3.column_dimensions["B"].width = 110

    # ── Notes sheet (from latest draft) ──
    if latest_draft:
        pinned = latest_draft.get("selected_insights") or []
        notes_text = (latest_draft.get("notes") or "").strip()
        if pinned or notes_text:
            s4 = wb.create_sheet("Notes")
            row = 1
            if pinned:
                s4.cell(row=row, column=1, value="Pinned for decision").font = header_font
                s4.cell(row=row, column=1).fill = header_fill
                row += 1
                for p_text in pinned:
                    cell = s4.cell(row=row, column=1, value=f"• {p_text}")
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
                    row += 1
                row += 1
            if notes_text:
                s4.cell(row=row, column=1, value="Analyst notes").font = header_font
                s4.cell(row=row, column=1).fill = header_fill
                row += 1
                for line in notes_text.split("\n"):
                    if line.strip():
                        cell = s4.cell(row=row, column=1, value=line.strip())
                        cell.alignment = Alignment(wrap_text=True, vertical="top")
                        row += 1
            s4.column_dimensions["A"].width = 110

    wb.save(str(output_path))
    return output_path
